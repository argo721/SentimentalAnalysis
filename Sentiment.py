from google.cloud import language
from google.cloud.language import enums, types
import pandas as pd
import openpyxl
import os
from matplotlib import pyplot


os.environ["GOOGLE_APPLICATION_CREDENTIALS"]= "C:\Hult\Internship\My First Project-6fc16db2f85f.json"
print(os.environ['GOOGLE_APPLICATION_CREDENTIALS'])

def sentiment_analysis(text):
    client = language.LanguageServiceClient()

    document = types.Document(
        content=text,
        type=enums.Document.Type.PLAIN_TEXT)

    sent_analysis = client.analyze_sentiment(document = document)
    dir(sent_analysis)
    sentiment = sent_analysis.document_sentiment

    return sentiment


def entity_analysis(text):
    client = language.LanguageServiceClient()

    document = types.Document(
        content=text,
        type=enums.Document.Type.PLAIN_TEXT)

    ent_analysis = client.analyze_entities(document=document)
    dir(ent_analysis)
    entities = ent_analysis.entities

    return entities



excel_file_path="C:\Hult\Internship\SentimentAnalysis_Lessons.xlsx"

sentiment_df = pd.read_excel(excel_file_path)
print(sentiment_df.head())
print(len(sentiment_df.Actual_Text))

sentimental_excel = openpyxl.open(excel_file_path)
sheet = sentimental_excel.active


#Creating an empty arrays for correlation  analysis
len_arr = []
mag_arr = []

for i in range(len(sentiment_df.Actual_Text)):
    print(sentiment_df.Actual_Text[i])
    print(len(sentiment_df.Actual_Text[i]))
    sentiment = sentiment_analysis(sentiment_df.Actual_Text[i])
    print("Score:" + str(sentiment.score) + " Magnitude:" + str(sentiment.magnitude) + "\n")
    sheet.cell(row = i+2, column = 5).value = sentiment.score
    sheet.cell(row = i+2, column = 6).value = sentiment.magnitude
    len_arr += [len(sentiment_df.Actual_Text[i])]
    mag_arr += [sentiment.magnitude]


sentimental_excel.save('SentimentAnalysis_Lessons.xlsx')

pyplot.scatter(len_arr, mag_arr)
pyplot.show()

example_text = 'Python is such a great programming language'
#sentiment = sentiment_analysis(example_text)
#print("Score:" + str(sentiment.score) + "Magnitude:" + sentiment.magnitude)


#entities =entity_analysise(text)
#for e in entities:
#    print(e.name, enums.Entity.Type(e.type).name, e.metadata, e.salience)



